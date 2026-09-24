"""
MarkdownToWord — محوّل Markdown إلى Word مع دعم RTL
====================================================

Copyright (c) 2025 عمادالدين الاديمي
GitHub:   https://github.com/emadaldin2003
Instagram: https://www.instagram.com/e.4_m

Licensed under the MIT License.
See LICENSE file for details.

الوضعان:
--------
- python md_to_docx_rtl.py                     → واجهة GUI
- python md_to_docx_rtl.py "file.md"           → تحويل مباشر
- python md_to_docx_rtl.py "file.md" "out.docx" → تحويل مع تحديد الإخراج

الميزات:
--------
- RTL/LTR صحيح مع إصلاح الأقواس
- جداول بتصميم Zebra + رأس أزرق داكن
- اقتباسات بشريط جانبي أزرق
- أكواد بإطار ناعم + inline code ملون
- عناوين ملونة + خط سفلي لـ H1/H2
- ترقيم صفحات في التذييل
- فهرس تلقائي عند وجود "المحتويات"/"الفهرس"
- حفظ آمن (اسم بديل عند فشل الكتابة)
- واجهة GUI عربية بسيطة
"""

import os
import re
import sys
from pathlib import Path
from typing import Union

import markdown
from bs4 import BeautifulSoup
from bs4.element import NavigableString, Tag

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Pt, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


# ============================================================
# الإعدادات
# ============================================================

DEFAULT_FONT = "Arial"
CODE_FONT = "Consolas"

BODY_SIZE = 18
CODE_SIZE = 12

HEADING_SIZES = {1: 28, 2: 24, 3: 21, 4: 19, 5: 18, 6: 18}

AVAILABLE_WIDTH_CM = 17.0

TABLE_FONT_SIZES = {
    1: 15, 2: 14, 3: 13, 4: 13,
    5: 12, 6: 11, 7: 10, 8: 10,
}


# ============================================================
# لوحة الألوان
# ============================================================

COLORS = {
    "h1": "1A365D",
    "h2": "2C5282",
    "h3": "2D3748",
    "h4": "4A5568",
    "h5": "4A5568",
    "h6": "4A5568",
    "bold": "2B6CB0",
    "body": "1A202C",
    "table_header_bg": "2C5282",
    "table_header_fg": "FFFFFF",
    "table_stripe": "F7FAFC",
    "table_border": "CBD5E0",
    "blockquote_bg": "EDF2F7",
    "blockquote_border": "4299E1",
    "code_bg": "F7FAFC",
    "code_border": "E2E8F0",
    "inline_code_bg": "FFF5F5",
    "inline_code_fg": "C53030",
    "code_fg": "2D3748",
    "hr": "A0AEC0",
    "footer": "A0AEC0",
}


# ============================================================
# ترتيب عناصر XML حسب OOXML Schema
# ============================================================

_RPR_ORDER = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
    "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
    "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
    "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
    "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
]

_PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore",
    "w:framePr", "w:widowControl", "w:numPr", "w:suppressLineNumbers",
    "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
    "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
    "w:rPr", "w:sectPr", "w:pPrChange",
]

_TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge",
    "w:tcBorders", "w:shd", "w:noWrap", "w:tcMar", "w:textDirection",
    "w:tcFitText", "w:vAlign", "w:hideMark",
]


def _insert_in_order(parent, element, tag_name, order):
    existing = parent.find(qn(tag_name))
    if existing is not None:
        return existing
    try:
        idx = order.index(tag_name)
    except ValueError:
        parent.append(element)
        return element
    for child in parent:
        child_name = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        child_full = 'w:' + child_name
        try:
            child_idx = order.index(child_full)
            if child_idx > idx:
                child.addprevious(element)
                return element
        except ValueError:
            continue
    parent.append(element)
    return element


def rpr_add(r_pr, tag_name):
    return _insert_in_order(r_pr, OxmlElement(tag_name), tag_name, _RPR_ORDER)


def ppr_add(p_pr, tag_name):
    return _insert_in_order(p_pr, OxmlElement(tag_name), tag_name, _PPR_ORDER)


def tcpr_add(tc_pr, tag_name):
    return _insert_in_order(tc_pr, OxmlElement(tag_name), tag_name, _TCPR_ORDER)


# ============================================================
# دالة محورية: ضبط حجم الخط للعربية واللاتينية
# ============================================================

def set_run_size(run, size_pt):
    """
    ضبط w:sz (للنص العادي) و w:szCs (للنص العربي).

    Word يستخدم w:szCs للعربية، وإذا لم يُضبط يستعمل حجماً
    افتراضياً (~11pt) بغض النظر عن w:sz.
    """
    r_pr = run._r.get_or_add_rPr()
    half_points = str(int(round(size_pt * 2)))

    sz = r_pr.find(qn("w:sz"))
    if sz is None:
        sz = rpr_add(r_pr, "w:sz")
    sz.set(qn("w:val"), half_points)

    sz_cs = r_pr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = rpr_add(r_pr, "w:szCs")
    sz_cs.set(qn("w:val"), half_points)

    run.font.size = Pt(size_pt)


# ============================================================
# كشف العربية واللاتينية
# ============================================================

ARABIC_RE = re.compile(r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF]")
LATIN_RE = re.compile(r"[A-Za-z]")


def contains_arabic(s):
    return bool(ARABIC_RE.search(s))


def contains_latin(s):
    return bool(LATIN_RE.search(s))


# ============================================================
# تقسيم النص حسب الاتجاه
# ============================================================

TOKEN_RE = re.compile(
    r"("
    r"\([^()]*\)"
    r"|\[[^\[\]]*\]"
    r"|\{[^{}]*\}"
    r"|«[^»]*»"
    r"|\u201C[^\u201D]*\u201D"
    r"|\"[^\"]*\""
    r"|\u2018[^\u2019]*\u2019"
    r"|'[^']*'"
    r"|https?://\S+"
    r"|[A-Za-z][A-Za-z0-9_.\-]*"
    r"|\d+(?:[.,:/-]\d+)*"
    r")"
)


def get_direction(text):
    if contains_latin(text) and not contains_arabic(text):
        return "ltr"
    return "rtl"


def split_direction(text):
    if not text:
        return []
    result = []
    last_end = 0
    for m in TOKEN_RE.finditer(text):
        start, end = m.span()
        if start > last_end:
            between = text[last_end:start]
            if between:
                d = get_direction(between)
                if result and result[-1][0] == d:
                    result[-1] = (d, result[-1][1] + between)
                else:
                    result.append((d, between))
        token = m.group(0)
        d = get_direction(token)
        if result and result[-1][0] == d:
            result[-1] = (d, result[-1][1] + token)
        else:
            result.append((d, token))
        last_end = end
    if last_end < len(text):
        trailing = text[last_end:]
        if trailing:
            d = get_direction(trailing)
            if result and result[-1][0] == d:
                result[-1] = (d, result[-1][1] + trailing)
            else:
                result.append((d, trailing))
    return result


# ============================================================
# المعالجة المسبقة للـ Markdown
# ============================================================

TABLE_ROW_RE = re.compile(r"^\s*\|.*\|\s*$")
LIST_ITEM_RE = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+")
HEADING_RE = re.compile(r"^\s*#{1,6}\s")
CODE_FENCE_RE = re.compile(r"^\s*```")


def _strip_trailing_backslash(lines):
    result = list(lines)
    for i in range(len(result) - 1):
        line = result[i]
        if not line.rstrip().endswith("\\"):
            continue
        nxt = result[i + 1]
        if LIST_ITEM_RE.match(nxt) or TABLE_ROW_RE.match(nxt):
            result[i] = line.rstrip()[:-1].rstrip()
    return result


def preprocess_markdown(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    lines = _strip_trailing_backslash(lines)

    in_code = False
    kinds = []
    for line in lines:
        s = line.strip()
        if CODE_FENCE_RE.match(line):
            kinds.append("code_fence")
            in_code = not in_code
            continue
        if in_code:
            kinds.append("code")
            continue
        if s == "":
            kinds.append("blank")
        elif TABLE_ROW_RE.match(line):
            kinds.append("table")
        elif LIST_ITEM_RE.match(line):
            kinds.append("list")
        elif HEADING_RE.match(line):
            kinds.append("heading")
        else:
            kinds.append("text")

    result = []
    for i, line in enumerate(lines):
        cur = kinds[i]
        if i > 0:
            prev = kinds[i - 1]
            need = False
            if cur == "list" and prev not in ("list", "blank", "code", "code_fence"):
                need = True
            elif cur == "table" and prev not in ("table", "blank", "code", "code_fence"):
                need = True
            elif prev == "table" and cur not in ("table", "blank"):
                need = True
            elif cur == "code_fence" and prev != "blank":
                need = True
            if need:
                result.append("")
        result.append(line)
    return "\n".join(result)


# ============================================================
# الفقرات والأنماط
# ============================================================

def set_paragraph_direction(paragraph, rtl=True):
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = ppr_add(p_pr, "w:bidi")
    bidi.set(qn("w:val"), "1" if rtl else "0")
    jc = ppr_add(p_pr, "w:jc")
    jc.set(qn("w:val"), "start" if rtl else "left")


def apply_line_spacing(paragraph, line=1.5, before_pt=6, after_pt=8):
    p_pr = paragraph._p.get_or_add_pPr()
    spacing = ppr_add(p_pr, "w:spacing")
    spacing.set(qn("w:line"), str(int(line * 240)))
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:before"), str(int(before_pt * 20)))
    spacing.set(qn("w:after"), str(int(after_pt * 20)))


def apply_paragraph_shading(paragraph, fill_hex):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = ppr_add(p_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def apply_paragraph_border(paragraph, sides):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = ppr_add(p_pr, "w:pBdr")
    for old in list(p_bdr):
        p_bdr.remove(old)
    order = ["top", "left", "bottom", "right"]
    for side_name in order:
        if side_name not in sides:
            continue
        style, sz, color = sides[side_name]
        el = OxmlElement(f"w:{side_name}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), color)
        p_bdr.append(el)


def apply_run_color(run, color_hex):
    r_pr = run._r.get_or_add_rPr()
    existing = r_pr.find(qn("w:color"))
    if existing is not None:
        existing.set(qn("w:val"), color_hex)
        return
    c = rpr_add(r_pr, "w:color")
    c.set(qn("w:val"), color_hex)


def apply_run_shading(run, fill_hex):
    r_pr = run._r.get_or_add_rPr()
    shd = rpr_add(r_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


def apply_cell_shading(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tcpr_add(tc_pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)


# ============================================================
# الخط
# ============================================================

def configure_run(run, font=DEFAULT_FONT, size=BODY_SIZE):
    run.font.name = font
    r_pr = run._r.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = rpr_add(r_pr, "w:rFonts")
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    r_fonts.set(qn("w:cs"), font)
    set_run_size(run, size)


def set_run_ltr(run):
    r_pr = run._r.get_or_add_rPr()
    rtl = rpr_add(r_pr, "w:rtl")
    rtl.set(qn("w:val"), "0")


def set_run_rtl(run):
    r_pr = run._r.get_or_add_rPr()
    rtl = rpr_add(r_pr, "w:rtl")
    rtl.set(qn("w:val"), "1")


# ============================================================
# نص عادي
# ============================================================

def add_mixed_text(paragraph, text, bold=False, italic=False, color=None, size=BODY_SIZE):
    if not text:
        return
    parts = split_direction(text)
    for direction, value in parts:
        if not value:
            continue
        run = paragraph.add_run(value)
        configure_run(run, size=size)
        run.bold = bold
        run.italic = italic
        if direction == "ltr":
            set_run_ltr(run)
        else:
            set_run_rtl(run)
        if color:
            apply_run_color(run, color)


# ============================================================
# Hyperlink
# ============================================================

def add_hyperlink(paragraph, url, text):
    try:
        r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    except Exception:
        add_mixed_text(paragraph, text)
        return

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")

    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), DEFAULT_FONT)
    r_fonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    r_fonts.set(qn("w:cs"), DEFAULT_FONT)
    r_pr.append(r_fonts)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2B6CB0")
    r_pr.append(color)

    hp = str(BODY_SIZE * 2)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), hp); r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs"); sz_cs.set(qn("w:val"), hp); r_pr.append(sz_cs)

    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); r_pr.append(u)
    rtl = OxmlElement("w:rtl"); rtl.set(qn("w:val"), "0"); r_pr.append(rtl)

    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


# ============================================================
# معالجة العناصر السطرية
# ============================================================

def process_inline(element, paragraph, bold=False, italic=False, bold_color=None):
    for child in element.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text:
                color = bold_color if bold else None
                add_mixed_text(paragraph, text, bold=bold, italic=italic, color=color)
            continue

        if not isinstance(child, Tag):
            continue

        tag = (child.name or "").lower()
        if not tag:
            continue

        if tag in ("strong", "b"):
            process_inline(child, paragraph, bold=True, italic=italic,
                           bold_color=COLORS["bold"])
        elif tag in ("em", "i"):
            process_inline(child, paragraph, bold=bold, italic=True,
                           bold_color=bold_color)
        elif tag == "code":
            run = paragraph.add_run(child.get_text())
            configure_run(run, font=CODE_FONT, size=CODE_SIZE)
            apply_run_color(run, COLORS["inline_code_fg"])
            apply_run_shading(run, COLORS["inline_code_bg"])
            set_run_ltr(run)
        elif tag == "a":
            url_value = child.get("href", "")
            if isinstance(url_value, list):
                url = str(url_value[0]) if url_value else ""
            else:
                url = str(url_value) if url_value is not None else ""
            text = child.get_text()
            if url:
                add_hyperlink(paragraph, url, text)
            else:
                add_mixed_text(paragraph, text)
        elif tag == "span":
            process_inline(child, paragraph, bold=bold, italic=italic,
                           bold_color=bold_color)
        elif tag == "br":
            paragraph.add_run().add_break()
        elif tag in ("del", "s", "strike"):
            start = len(paragraph.runs)
            process_inline(child, paragraph, bold=bold, italic=italic,
                           bold_color=bold_color)
            for run in paragraph.runs[start:]:
                run.font.strike = True
        elif tag == "img":
            alt_value = child.get("alt", "")
            if isinstance(alt_value, list):
                alt = str(alt_value[0]) if alt_value else ""
            else:
                alt = str(alt_value) if alt_value is not None else ""
            if alt:
                add_mixed_text(paragraph, f"[صورة: {alt}]", color=COLORS["h4"])
        else:
            process_inline(child, paragraph, bold=bold, italic=italic,
                           bold_color=bold_color)


# ============================================================
# فقرة
# ============================================================

def add_paragraph(doc, element):
    paragraph = doc.add_paragraph()
    set_paragraph_direction(paragraph, rtl=True)
    process_inline(element, paragraph)
    apply_line_spacing(paragraph, line=1.5, before_pt=4, after_pt=8)

    for run in paragraph.runs:
        if run.font.size is None:
            set_run_size(run, BODY_SIZE)
        apply_run_color(run, COLORS["body"])

    return paragraph


# ============================================================
# عنوان
# ============================================================

def _is_toc_heading(text):
    t = text.strip().lower()
    return any(k in t for k in ["المحتويات", "الفهرس", "فهرس", "table of contents", "contents"])


def add_heading(doc, element, level, add_toc_after=False):
    word_level = max(1, min(level, 9))
    paragraph = doc.add_heading(level=word_level)
    set_paragraph_direction(paragraph, rtl=True)
    process_inline(element, paragraph)

    size = HEADING_SIZES.get(level, BODY_SIZE)
    color = COLORS.get(f"h{min(level, 6)}", COLORS["h4"])

    for run in paragraph.runs:
        set_run_size(run, size)
        apply_run_color(run, color)

    if level == 1:
        apply_paragraph_border(paragraph, {"bottom": ("single", 12, COLORS["h1"])})
    elif level == 2:
        apply_paragraph_border(paragraph, {"bottom": ("single", 8, COLORS["h2"])})

    before = max(8, 18 - level * 2)
    after = max(4, 10 - level)
    apply_line_spacing(paragraph, line=1.15, before_pt=before, after_pt=after)

    if add_toc_after:
        _insert_toc_field(doc)

    return paragraph


# ============================================================
# TOC
# ============================================================

def _insert_toc_field(doc):
    paragraph = doc.add_paragraph()
    set_paragraph_direction(paragraph, rtl=True)

    r1 = paragraph.add_run()
    fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc1)

    r2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    r2._r.append(instr)

    r3 = paragraph.add_run()
    fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate")
    r3._r.append(fc2)

    r4 = paragraph.add_run("اضغط Ctrl+A ثم F9 لتحديث الفهرس")
    configure_run(r4, size=11)
    apply_run_color(r4, COLORS["h4"])
    r4.italic = True

    r5 = paragraph.add_run()
    fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end")
    r5._r.append(fc3)


# ============================================================
# Code Block
# ============================================================

def add_code_block(doc, element):
    code_el = element.find("code")
    source = code_el if code_el is not None else element
    code_text = source.get_text().rstrip("\n")
    if not code_text:
        return None

    paragraph = doc.add_paragraph()
    set_paragraph_direction(paragraph, rtl=False)

    apply_paragraph_border(paragraph, {
        "top": ("single", 4, COLORS["code_border"]),
        "bottom": ("single", 4, COLORS["code_border"]),
        "left": ("single", 4, COLORS["code_border"]),
        "right": ("single", 4, COLORS["code_border"]),
    })
    apply_paragraph_shading(paragraph, COLORS["code_bg"])

    p_pr = paragraph._p.get_or_add_pPr()
    ind = ppr_add(p_pr, "w:ind")
    ind.set(qn("w:left"), "120")
    ind.set(qn("w:right"), "120")

    apply_line_spacing(paragraph, line=1.15, before_pt=6, after_pt=8)

    lines = code_text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        configure_run(run, font=CODE_FONT, size=CODE_SIZE)
        apply_run_color(run, COLORS["code_fg"])
        set_run_ltr(run)

    return paragraph


# ============================================================
# قائمة
# ============================================================

def add_list(doc, element, ordered=False, level=0):
    items = element.find_all("li", recursive=False)
    for item in items:
        nested = []
        for child in list(item.children):
            if isinstance(child, Tag):
                ct = (child.name or "").lower()
                if ct in ("ul", "ol"):
                    nested.append(child)

        if ordered:
            style = "List Number" if level == 0 else f"List Number {min(level + 1, 3)}"
        else:
            style = "List Bullet" if level == 0 else f"List Bullet {min(level + 1, 3)}"

        try:
            paragraph = doc.add_paragraph(style=style)
        except KeyError:
            paragraph = doc.add_paragraph(style="List Bullet" if not ordered else "List Number")

        set_paragraph_direction(paragraph, rtl=True)

        for child in item.children:
            if isinstance(child, Tag):
                ct = (child.name or "").lower()
                if ct in ("ul", "ol"):
                    continue
            if isinstance(child, NavigableString):
                t = str(child)
                if t.strip():
                    add_mixed_text(paragraph, t)
            elif isinstance(child, Tag):
                process_inline(child, paragraph)

        for run in paragraph.runs:
            if run.font.size is None:
                set_run_size(run, BODY_SIZE)
            apply_run_color(run, COLORS["body"])

        apply_line_spacing(paragraph, line=1.4, before_pt=2, after_pt=4)

        for n in nested:
            nt = (n.name or "").lower()
            add_list(doc, n, ordered=(nt == "ol"), level=level + 1)


# ============================================================
# اقتباس
# ============================================================

def add_blockquote(doc, element):
    paragraph = doc.add_paragraph()
    set_paragraph_direction(paragraph, rtl=True)

    apply_paragraph_border(paragraph, {"right": ("single", 18, COLORS["blockquote_border"])})
    apply_paragraph_shading(paragraph, COLORS["blockquote_bg"])

    p_pr = paragraph._p.get_or_add_pPr()
    ind = ppr_add(p_pr, "w:ind")
    ind.set(qn("w:right"), "200")
    ind.set(qn("w:left"), "200")

    process_inline(element, paragraph)
    for run in paragraph.runs:
        set_run_size(run, BODY_SIZE)
        apply_run_color(run, "2D3748")
        run.italic = True

    apply_line_spacing(paragraph, line=1.5, before_pt=6, after_pt=8)
    return paragraph


# ============================================================
# الجداول
# ============================================================

def set_table_rtl(table):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:bidiVisual")):
        tbl_pr.remove(old)
    b = OxmlElement("w:bidiVisual")
    tbl_pr.append(b)


def set_repeat_header(row):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        tr_pr.append(h)


def set_table_borders(table, color_hex):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color_hex)
        b.append(el)
    tbl_pr.append(b)


def set_table_cell_margins(table, top=60, left=100, bottom=60, right=100):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblCellMar")):
        tbl_pr.remove(old)
    cm = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        cm.append(el)
    tbl_pr.append(cm)


def compute_column_widths(rows, max_columns, available_cm=AVAILABLE_WIDTH_CM):
    col_weights = [0.0] * max_columns
    for row in rows:
        cells = row.find_all(["th", "td"])
        for i, cell in enumerate(cells):
            if i >= max_columns:
                break
            text = cell.get_text().strip()
            if not text:
                continue
            words = text.split()
            longest = max((len(w) for w in words), default=0)
            text_w = len(text) / 20.0
            col_weights[i] = max(col_weights[i], max(longest, text_w))
    MIN_W = 4.0
    col_weights = [max(w, MIN_W) for w in col_weights]
    total = sum(col_weights) or 1.0
    widths = [(w / total) * available_cm for w in col_weights]
    MIN_CM, MAX_CM = 1.5, 8.0
    widths = [max(MIN_CM, min(w, MAX_CM)) for w in widths]
    tw = sum(widths)
    if tw > available_cm:
        s = available_cm / tw
        widths = [w * s for w in widths]
    return widths


def set_grid_columns(table, widths_cm):
    tbl = table._tbl
    tg = tbl.find(qn("w:tblGrid"))
    if tg is None:
        tg = OxmlElement("w:tblGrid")
        tbl_pr = tbl.find(qn("w:tblPr"))
        if tbl_pr is not None:
            tbl_pr.addnext(tg)
        else:
            tbl.insert(0, tg)
    for col in tg.findall(qn("w:gridCol")):
        tg.remove(col)
    for w_cm in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(Cm(w_cm).twips)))
        tg.append(col)


def configure_table_cell(cell, rtl=True):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        set_paragraph_direction(p, rtl=rtl)


def fill_table_cell(cell, html_cell, font_size, color_hex, bold=False):
    paragraphs_in_cell = html_cell.find_all("p", recursive=False)
    first_para = cell.paragraphs[0]
    set_paragraph_direction(first_para, rtl=True)

    def apply_style(para):
        apply_line_spacing(para, line=1.0, before_pt=0, after_pt=0)
        for run in para.runs:
            set_run_size(run, font_size)
            apply_run_color(run, color_hex)
            if bold:
                run.bold = True

    if len(paragraphs_in_cell) > 1:
        process_inline(paragraphs_in_cell[0], first_para)
        apply_style(first_para)
        for p in paragraphs_in_cell[1:]:
            new_p = cell.add_paragraph()
            set_paragraph_direction(new_p, rtl=True)
            process_inline(p, new_p)
            apply_style(new_p)
    else:
        process_inline(html_cell, first_para)
        apply_style(first_para)


def add_table(doc, element):
    rows = element.find_all("tr")
    if not rows:
        return None

    max_columns = 0
    for row in rows:
        cells = row.find_all(["th", "td"])
        max_columns = max(max_columns, len(cells))
    if max_columns == 0:
        return None

    font_size = TABLE_FONT_SIZES.get(max_columns, 10)

    table = doc.add_table(rows=len(rows), cols=max_columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = None

    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblW")):
        tbl_pr.remove(old)
    tw_el = OxmlElement("w:tblW")
    tw_el.set(qn("w:type"), "dxa")
    tw_el.set(qn("w:w"), str(int(Cm(AVAILABLE_WIDTH_CM).twips)))
    tbl_pr.append(tw_el)

    for old in tbl_pr.findall(qn("w:tblLayout")):
        tbl_pr.remove(old)
    tl = OxmlElement("w:tblLayout")
    tl.set(qn("w:type"), "fixed")
    tbl_pr.append(tl)

    set_table_rtl(table)
    set_table_borders(table, COLORS["table_border"])
    set_table_cell_margins(table)

    widths_cm = compute_column_widths(rows, max_columns)
    set_grid_columns(table, widths_cm)

    first_row_cells = rows[0].find_all(["th", "td"])
    has_header = any((c.name or "").lower() == "th" for c in first_row_cells)

    for r_idx, html_row in enumerate(rows):
        html_cells = html_row.find_all(["th", "td"])
        is_header_row = has_header and r_idx == 0

        if is_header_row:
            row_bg = COLORS["table_header_bg"]
            row_fg = COLORS["table_header_fg"]
            row_bold = True
        else:
            data_idx = r_idx - (1 if has_header else 0)
            row_bg = COLORS["table_stripe"] if data_idx % 2 == 1 else "FFFFFF"
            row_fg = COLORS["body"]
            row_bold = False

        for c_idx in range(max_columns):
            cell = table.cell(r_idx, c_idx)
            configure_table_cell(cell, rtl=True)
            cell.width = Cm(widths_cm[c_idx])
            apply_cell_shading(cell, row_bg)

            if c_idx >= len(html_cells):
                continue

            html_cell = html_cells[c_idx]
            is_th = (html_cell.name or "").lower() == "th"

            fill_table_cell(cell, html_cell, font_size, row_fg, bold=(row_bold or is_th))

    if has_header and rows:
        set_repeat_header(table.rows[0])

    spacer = doc.add_paragraph()
    apply_line_spacing(spacer, line=1.0, before_pt=0, after_pt=6)

    return table


# ============================================================
# خط أفقي
# ============================================================

def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    set_paragraph_direction(p, rtl=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("•    •    •")
    configure_run(run, size=14)
    apply_run_color(run, COLORS["hr"])
    apply_line_spacing(p, line=1.0, before_pt=6, after_pt=6)


# ============================================================
# تذييل الصفحة
# ============================================================

def add_page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False

    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for r in list(p.runs):
        r._r.getparent().remove(r._r)

    apply_paragraph_border(p, {"top": ("single", 4, COLORS["footer"])})

    run = p.add_run()
    configure_run(run, size=10)
    apply_run_color(run, COLORS["footer"])

    fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "end")
    run._r.append(fc1); run._r.append(instr); run._r.append(fc2)


# ============================================================
# إعداد المستند
# ============================================================

def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.right_margin = Cm(2)
    section.left_margin = Cm(2)

    sect_pr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sect_pr.append(bidi)

    normal = document.styles["Normal"]
    normal.font.name = DEFAULT_FONT
    normal.font.size = Pt(BODY_SIZE)

    rpr = normal.element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), DEFAULT_FONT)
    r_fonts.set(qn("w:hAnsi"), DEFAULT_FONT)
    r_fonts.set(qn("w:cs"), DEFAULT_FONT)

    hp = str(BODY_SIZE * 2)
    sz_el = rpr.find(qn("w:sz"))
    if sz_el is None:
        sz_el = OxmlElement("w:sz")
        rpr.append(sz_el)
    sz_el.set(qn("w:val"), hp)

    sz_cs = rpr.find(qn("w:szCs"))
    if sz_cs is None:
        sz_cs = OxmlElement("w:szCs")
        rpr.append(sz_cs)
    sz_cs.set(qn("w:val"), hp)

    add_page_number_footer(section)


# ============================================================
# Markdown → HTML
# ============================================================

def markdown_to_html(md):
    return markdown.markdown(md, extensions=["extra", "sane_lists"])


# ============================================================
# الحفظ الآمن
# ============================================================

def is_file_writable(path):
    if not path.exists():
        return True
    try:
        with open(path, "a"):
            pass
        return True
    except (PermissionError, OSError):
        return False


def save_document_safely(document, output_path):
    output_path = Path(output_path)
    try:
        document.save(str(output_path))
        return output_path
    except (PermissionError, OSError) as e:
        print()
        print(f"⚠️  تعذّر الكتابة فوق: {output_path.name}")
        print(f"   السبب: {e}")
        print("   سيتم استخدام اسم بديل...")
        print()

    stem, suffix, parent = output_path.stem, output_path.suffix, output_path.parent
    for i in range(1, 100):
        candidate = parent / f"{stem} ({i}){suffix}"
        try:
            document.save(str(candidate))
            print(f"✅ تم الحفظ باسم بديل: {candidate.name}")
            print()
            return candidate
        except (PermissionError, OSError):
            continue

    raise RuntimeError(f"تعذّر حفظ الملف بعد 100 محاولة.")


# ============================================================
# التحويل الرئيسي
# ============================================================

def convert_md_to_docx(input_file, output_file):
    input_path = Path(input_file)
    if not input_path.exists():
        raise FileNotFoundError(f"الملف غير موجود:\n{input_path}")

    md_text = input_path.read_text(encoding="utf-8")
    md_text = preprocess_markdown(md_text)
    html = markdown_to_html(md_text)
    soup = BeautifulSoup(html, "html.parser")

    document = Document()
    configure_document(document)

    def process_block(el):
        if isinstance(el, NavigableString):
            return
        if not isinstance(el, Tag):
            return
        tag = (el.name or "").lower()
        if not tag:
            return

        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            level = int(tag[1])
            text = el.get_text().strip()
            add_toc = (level <= 2) and _is_toc_heading(text)
            add_heading(document, el, level, add_toc_after=add_toc)
        elif tag == "p":
            add_paragraph(document, el)
        elif tag == "pre":
            add_code_block(document, el)
        elif tag == "table":
            add_table(document, el)
        elif tag == "ul":
            add_list(document, el, ordered=False)
        elif tag == "ol":
            add_list(document, el, ordered=True)
        elif tag == "blockquote":
            add_blockquote(document, el)
        elif tag == "hr":
            add_horizontal_rule(document)
        elif tag in ("div", "section", "article", "main"):
            for child in el.children:
                process_block(child)

    for el in soup.children:
        process_block(el)

    return save_document_safely(document, Path(output_file))


# ============================================================
# قسم GUI
# ============================================================

def run_gui():
    """
    تشغيل الواجهة الرسومية.
    تُستدعى عندما لا توجد وسائط سطر أوامر.
    """
    import threading
    import queue
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, simpledialog

    state: dict[str, str | None] = {
        "md_path": None,
        "out_path": None,
        "last_output": None,
    }
    # ---------------------------------------------------------
    # النافذة
    # ---------------------------------------------------------

    root = tk.Tk()
    root.title("محول Markdown إلى Word")
    root.geometry("760x480")
    root.minsize(660, 440)

    try:
        default_font = ("Segoe UI", 11)
    except Exception:
        default_font = ("Arial", 11)

    root.option_add("*Font", default_font)
    root.configure(bg="#F7FAFC")

    style = ttk.Style()
    try:
        style.theme_use("clam")
    except Exception:
        pass

    style.configure("TFrame", background="#F7FAFC")
    style.configure("TLabel", background="#F7FAFC", foreground="#1A202C")
    style.configure("Title.TLabel",
                    font=("Segoe UI", 18, "bold"),
                    foreground="#1A365D",
                    background="#F7FAFC")
    style.configure("Sub.TLabel",
                    font=("Segoe UI", 10),
                    foreground="#4A5568",
                    background="#F7FAFC")
    style.configure("Status.TLabel",
                    font=("Segoe UI", 11, "bold"),
                    foreground="#2C5282",
                    background="#F7FAFC")
    style.configure("Primary.TButton",
                    font=("Segoe UI", 12, "bold"),
                    foreground="white",
                    background="#2C5282")
    style.map("Primary.TButton",
              background=[("active", "#2B6CB0"),
                          ("disabled", "#A0AEC0")])

    # ---------------------------------------------------------
    # التخطيط
    # ---------------------------------------------------------

    main_frame = ttk.Frame(root, padding=24)
    main_frame.pack(fill="both", expand=True)

    ttk.Label(
        main_frame,
        text="محول Markdown إلى Word",
        style="Title.TLabel",
        anchor="e",
        justify="right",
    ).pack(fill="x", pady=(0, 6))

    ttk.Label(
        main_frame,
        text="تحويل ملفات Markdown إلى Word مع الحفاظ على التنسيق واتجاه النص.",
        style="Sub.TLabel",
        anchor="e",
        justify="right",
    ).pack(fill="x", pady=(0, 20))

    ttk.Separator(main_frame, orient="horizontal").pack(fill="x", pady=(0, 20))

    # صف اختيار الملف
    file_frame = ttk.Frame(main_frame)
    file_frame.pack(fill="x", pady=(0, 12))

    ttk.Label(
        file_frame,
        text="ملف Markdown:",
        anchor="e",
    ).pack(side="right", padx=(0, 8))

    entry_path = ttk.Entry(file_frame, font=("Segoe UI", 10))
    entry_path.pack(side="right", fill="x", expand=True, padx=(0, 8))

    # صف أزرار التصفح
    browse_frame = ttk.Frame(main_frame)
    browse_frame.pack(fill="x", pady=(0, 16))

    browse_btn = ttk.Button(browse_frame, text="تصفح...")
    browse_btn.pack(side="right", padx=(0, 8))

    out_btn = ttk.Button(browse_frame, text="مكان الحفظ...")
    out_btn.pack(side="right")

    # زر التحويل
    convert_btn = ttk.Button(
        main_frame,
        text="تحويل إلى Word",
        style="Primary.TButton",
    )
    convert_btn.pack(fill="x", pady=(4, 16), ipady=8)

    # الحالة
    status_frame = ttk.Frame(main_frame)
    status_frame.pack(fill="x", pady=(0, 16))

    ttk.Label(status_frame, text="الحالة:", anchor="e").pack(side="right", padx=(0, 8))

    status_label = ttk.Label(
        status_frame,
        text="جاهز",
        style="Status.TLabel",
        anchor="e",
    )
    status_label.pack(side="right", fill="x", expand=True)

    # زر فتح الملف
    open_btn = ttk.Button(main_frame, text="فتح الملف")
    open_btn.pack(fill="x", ipady=6)
    open_btn.state(["disabled"])

    # ---------------------------------------------------------
    # دوال مساعدة
    # ---------------------------------------------------------

    def set_status(text, color="#2C5282"):
        status_label.configure(text=text, foreground=color)

    def browse_file():
        chosen = filedialog.askopenfilename(
            title="اختر ملف Markdown",
            filetypes=[
                ("ملفات Markdown", "*.md"),
                ("ملفات نصية", "*.markdown *.txt"),
                ("كل الملفات", "*.*"),
            ],
        )
        if not chosen:
            return

        p = Path(chosen)
        if p.suffix.lower() not in (".md", ".markdown", ".txt"):
            ans = messagebox.askyesno(
                "تنبيه",
                "الملف المختار ليس ملف Markdown بامتداد .md\n"
                "هل تريد المتابعة على أي حال؟",
            )
            if not ans:
                return

        state["md_path"] = str(p)
        entry_path.delete(0, "end")
        entry_path.insert(0, str(p))
        state["last_output"] = None
        open_btn.state(["disabled"])
        set_status("جاهز — الملف محدد")

    def browse_output():
        if not state["md_path"]:
            messagebox.showwarning("تنبيه", "اختر ملف Markdown أولًا.")
            return
        base = Path(state["md_path"]).with_suffix(".docx")
        chosen = filedialog.asksaveasfilename(
            title="اختر مكان الحفظ",
            defaultextension=".docx",
            initialfile=base.name,
            initialdir=str(base.parent),
            filetypes=[("ملفات Word", "*.docx")],
        )
        if chosen:
            state["out_path"] = chosen
            set_status(f"سيُحفظ باسم: {Path(chosen).name}")

    def open_result():
        if not state["last_output"]:
            return
        try:
            os.startfile(str(state["last_output"]))
        except AttributeError:
            import subprocess
            if sys.platform == "darwin":
                subprocess.Popen(["open", str(state["last_output"])])
            else:
                subprocess.Popen(["xdg-open", str(state["last_output"])])
        except Exception as e:
            messagebox.showerror("خطأ", f"تعذّر فتح الملف:\n{e}")

    # ---------------------------------------------------------
    # التحويل في Thread
    # ---------------------------------------------------------

    def conversion_worker(md_path, out_path, q):
        try:
            actual = convert_md_to_docx(md_path, out_path)
            q.put(("ok", str(actual)))
        except Exception as e:
            q.put(("error", str(e)))

    def do_convert():
        md = state["md_path"] or entry_path.get().strip()

        if not md:
            messagebox.showwarning("تنبيه", "الرجاء اختيار ملف Markdown أولًا.")
            return

        md_path = Path(md)
        if not md_path.exists():
            messagebox.showerror("خطأ", f"الملف غير موجود:\n{md_path}")
            return

        if md_path.suffix.lower() not in (".md", ".markdown", ".txt"):
            ans = messagebox.askyesno(
                "تنبيه",
                "الملف ليس بامتداد .md\nهل تريد المتابعة؟",
            )
            if not ans:
                return

        try:
            if md_path.stat().st_size == 0:
                messagebox.showwarning("تنبيه", "الملف فارغ.")
                return
        except OSError as e:
            messagebox.showerror("خطأ", f"لا يمكن قراءة الملف:\n{e}")
            return

        out_path = state["out_path"]
        if not out_path:
            out_path = str(md_path.with_suffix(".docx"))

        out_p = Path(out_path)
        if out_p.exists():
            ans = messagebox.askyesno(
                "الملف موجود",
                f"الملف التالي موجود بالفعل:\n{out_p.name}\n\nهل تريد استبداله؟",
            )
            if not ans:
                new_name = simpledialog.askstring(
                    "اسم بديل",
                    "اكتب اسم ملف جديد (بدون .docx):",
                    initialvalue=out_p.stem + " - نسخة",
                )
                if not new_name:
                    return
                out_path = str(out_p.with_name(new_name + ".docx"))
                state["out_path"] = out_path
                out_p = Path(out_path)

        if out_p.exists() and not is_file_writable(out_p):
            messagebox.showinfo(
                "ملاحظة",
                "الملف الناتج مفتوح في برنامج آخر.\n"
                "سيتم إنشاء ملف باسم بديل تلقائيًا.",
            )

        convert_btn.state(["disabled"])
        browse_btn.state(["disabled"])
        out_btn.state(["disabled"])
        open_btn.state(["disabled"])
        set_status("جاري التحويل...", color="#DD6B20")

        q = queue.Queue()
        t = threading.Thread(
            target=conversion_worker,
            args=(str(md_path), out_path, q),
            daemon=True,
        )
        t.start()

        def poll():
            try:
                kind, payload = q.get_nowait()
            except queue.Empty:
                root.after(150, poll)
                return

            convert_btn.state(["!disabled"])
            browse_btn.state(["!disabled"])
            out_btn.state(["!disabled"])

            if kind == "ok":
                state["last_output"] = payload
                state["out_path"] = payload
                set_status(f"تم التحويل بنجاح: {Path(payload).name}", color="#2F855A")
                open_btn.state(["!disabled"])
                messagebox.showinfo(
                    "نجاح",
                    f"تم التحويل بنجاح.\n\nالملف الناتج:\n{payload}",
                )
            else:
                set_status("فشل التحويل", color="#C53030")
                messagebox.showerror(
                    "خطأ أثناء التحويل",
                    f"تعذّر إتمام التحويل:\n\n{payload}",
                )

        root.after(150, poll)

    # ربط الأزرار
    browse_btn.configure(command=browse_file)
    out_btn.configure(command=browse_output)
    convert_btn.configure(command=do_convert)
    open_btn.configure(command=open_result)

    root.bind("<Return>", lambda e: do_convert())
    
        # حقوق النشر في الأسفل (خارج main_frame)
    footer_label = ttk.Label(
        root,              # ← child of root مباشرة
        text="© 2025 عمادالدين الاديمي — MIT License",
        style="Sub.TLabel",
        anchor="center",
    )
    footer_label.pack(fill="x", side="bottom", pady=(0, 8))

    root.mainloop()


# ============================================================
# Main
# ============================================================

def main():
    # وضع GUI: لا توجد وسائط
    if len(sys.argv) < 2:
        try:
            run_gui()
        except Exception as e:
            try:
                import tkinter.messagebox as mb
                mb.showerror("خطأ فادح", str(e))
            except Exception:
                print(f"حدث خطأ: {e}", file=sys.stderr)
        return

    # وضع CLI
    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) >= 3 else str(Path(input_file).with_suffix(".docx"))

    if not is_file_writable(Path(output_file)):
        print()
        print("⚠️  تحذير مبكّر:")
        print(f"   الملف '{output_file}' مفتوح في برنامج آخر.")
        print("   سيتم إنشاء ملف باسم بديل عند الحفظ.")
        print()

    try:
        print()
        print("جاري تحويل الملف...")
        print()
        actual_path = convert_md_to_docx(input_file, output_file)
        print()
        print("تم التحويل بنجاح.")
        print()
        print("الملف الناتج:")
        print(Path(actual_path).resolve())
        print()
    except Exception as error:
        print()
        print("حدث خطأ أثناء التحويل:")
        print()
        print(error)
        print()
        import traceback
        traceback.print_exc()


# ============================================================
# Program entry point
# ============================================================

if __name__ == "__main__":
    main()